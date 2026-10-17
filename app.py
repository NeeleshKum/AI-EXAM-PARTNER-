"""ExamForge - turn any document into a competitive-exam mock test.

Stateless Flask backend: nothing is stored server-side, so a shared link works
for everyone independently. The API key never leaves the server.
"""
import io
import json
import os
import random
import re
import time
from collections import defaultdict, deque

import requests
from docx import Document
from dotenv import load_dotenv
from flask import Flask, jsonify, render_template, request

load_dotenv()

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
MAX_UPLOAD_MB = int(os.getenv("MAX_UPLOAD_MB", "15"))
MAX_CHARS = int(os.getenv("MAX_CHARS", "90000"))      # text sent to the model
MIN_CHARS = 300                                       # below this, too little to test on
RATE_LIMIT = int(os.getenv("RATE_LIMIT_PER_HOUR", "12"))
ALLOWED = {"pdf", "docx", "txt", "md"}

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_MB * 1024 * 1024

_hits = defaultdict(deque)


class UserError(Exception):
    def __init__(self, msg, status=400):
        super().__init__(msg)
        self.status = status


# --------------------------------------------------------------------------- #
# Text extraction
# --------------------------------------------------------------------------- #
def extract_pdf(data: bytes) -> str:
    text = ""
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text = "\n\n".join((p.extract_text() or "") for p in pdf.pages)
    except Exception:
        text = ""
    if len(text.strip()) < MIN_CHARS:  # fallback parser
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise UserError("This PDF is password-protected. Remove the password and try again.")
            text = "\n\n".join((p.extract_text() or "") for p in reader.pages)
        except UserError:
            raise
        except Exception:
            raise UserError("Couldn't read this PDF. It may be corrupted.")
    return text


def extract_docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        raise UserError("Couldn't read this Word file. Make sure it's a valid .docx (not .doc).")
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def extract_plain(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    raise UserError("Couldn't decode this text file.")


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED:
        raise UserError("Unsupported file type. Upload a PDF, DOCX, TXT or MD file.")
    text = {"pdf": extract_pdf, "docx": extract_docx}.get(ext, extract_plain)(data)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < MIN_CHARS:
        raise UserError(
            "Very little text found. If this is a scanned PDF, run OCR on it first.")
    return text


def fit_text(text: str):
    """Keep within the model budget. Long docs: sample evenly across the whole
    document so every section is represented, not just the opening pages."""
    if len(text) <= MAX_CHARS:
        return text, False
    chunks, n = [], 12
    size = MAX_CHARS // n
    step = len(text) // n
    for i in range(n):
        chunks.append(text[i * step: i * step + size])
    return "\n\n[...]\n\n".join(chunks), True


# --------------------------------------------------------------------------- #
# LLM
# --------------------------------------------------------------------------- #
PROMPT = """You are a senior question-setter for Indian competitive exams (SSC CGL Tier-2, \
UPSC Prelims). Using ONLY the source document below, write {n} multiple-choice questions \
at {level} standard.

Quality rules:
- Test analytical reasoning, conceptual clarity, inference, cause-effect, data interpretation \
or application. Avoid trivial word-matching or "which sentence appears in the text" questions.
- Mix formats: statement-based ("Which of the statements is/are correct?"), assertion-reason, \
match/sequence, "except"-type, numerical/data based (only if the document has data), and direct conceptual.
- Spread questions across the whole document and across its different sub-topics.
- Every question has EXACTLY 5 options: 1 correct and 4 plausible distractors built from \
common misconceptions or near-miss facts. Never use "All of the above" or "None of the above".
- Options must be similar in length and style; the correct answer must be unambiguous \
and supported by the document.
- Write a short explanation (2-3 sentences) naming why the answer is right and why the strongest \
distractor is wrong.

Return ONLY JSON in this exact shape:
{{
  "title": "short title for this document",
  "topics": [
    {{"name": "sub-topic name", "key_points": ["3-5 concise revision notes"],
      "revision_tips": ["2-3 specific, actionable study tips for this sub-topic"]}}
  ],
  "questions": [
    {{"question": "...", "options": ["...", "...", "...", "...", "..."],
      "correct_index": 0, "topic": "must match one name in topics",
      "explanation": "..."}}
  ]
}}
Use 4-10 topics (the document's real sub-sections/themes).

SOURCE DOCUMENT:
\"\"\"
{doc}
\"\"\""""


def call_gemini(prompt: str) -> dict:
    if not GEMINI_API_KEY:
        raise UserError("Server is missing GEMINI_API_KEY. See the README.", 500)
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:generateContent"
    body = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {"temperature": 0.7, "responseMimeType": "application/json",
                             "maxOutputTokens": 32768},
    }
    try:
        r = requests.post(url, json=body, timeout=150,
                          headers={"x-goog-api-key": GEMINI_API_KEY})
    except requests.RequestException:
        raise UserError("Couldn't reach the AI service. Try again in a moment.", 502)
    if r.status_code == 429:
        raise UserError("The AI service is busy (rate limit). Wait a minute and retry.", 429)
    if r.status_code >= 400:
        raise UserError(f"AI service error ({r.status_code}). Check your API key and model name.", 502)
    try:
        raw = r.json()["candidates"][0]["content"]["parts"][0]["text"]
        raw = re.sub(r"^```(?:json)?|```$", "", raw.strip(), flags=re.M).strip()
        return json.loads(raw)
    except Exception:
        raise UserError("The AI returned an unreadable response. Please try again.", 502)


def clean(data: dict, n: int) -> dict:
    topics = []
    for t in data.get("topics", []):
        if isinstance(t, dict) and t.get("name"):
            topics.append({
                "name": str(t["name"]).strip(),
                "key_points": [str(x) for x in t.get("key_points", [])][:6],
                "revision_tips": [str(x) for x in t.get("revision_tips", [])][:4],
            })
    names = {t["name"] for t in topics}
    questions = []
    for q in data.get("questions", []):
        try:
            opts = [str(o).strip() for o in q["options"]]
            ci = int(q["correct_index"])
            # Strict rule: exactly 5 distinct options, exactly one correct.
            if len(opts) != 5 or len(set(o.lower() for o in opts)) != 5 or not 0 <= ci < 5:
                continue
            correct = opts[ci]
            random.shuffle(opts)  # remove the model's positional bias
            topic = str(q.get("topic", "General")).strip()
            if topic not in names:
                topics.append({"name": topic, "key_points": [], "revision_tips": []})
                names.add(topic)
            questions.append({
                "question": str(q["question"]).strip(),
                "options": opts,
                "correct_index": opts.index(correct),
                "topic": topic,
                "explanation": str(q.get("explanation", "")).strip(),
            })
        except (KeyError, ValueError, TypeError):
            continue
    return {"title": str(data.get("title", "Practice test"))[:120],
            "topics": topics, "questions": questions[:n]}


# --------------------------------------------------------------------------- #
# Routes
# --------------------------------------------------------------------------- #
def throttle(ip: str):
    now, q = time.time(), _hits[ip]
    while q and now - q[0] > 3600:
        q.popleft()
    if len(q) >= RATE_LIMIT:
        raise UserError("Hourly limit reached for this connection. Try again later.", 429)
    q.append(now)


@app.get("/")
def index():
    return render_template("index.html", max_mb=MAX_UPLOAD_MB)


@app.get("/health")
def health():
    return {"ok": True, "key_configured": bool(GEMINI_API_KEY)}


@app.post("/api/generate")
def generate():
    ip = (request.headers.get("X-Forwarded-For", request.remote_addr) or "").split(",")[0].strip()
    throttle(ip)
    f = request.files.get("file")
    if not f or not f.filename:
        raise UserError("Choose a file first.")
    try:
        n = max(5, min(30, int(request.form.get("count", 10))))
    except ValueError:
        n = 10
    level = {"cgl": "SSC CGL / CHSL Tier-2 (or higher)",
             "upsc": "UPSC Civil Services Prelims",
             "adv": "advanced, harder than SSC CGL Tier-2"}.get(request.form.get("level"), "SSC CGL Tier-2")

    text, truncated = fit_text(extract_text(f.filename, f.read()))
    # Ask for a few extra: invalid questions get dropped by validation.
    result = clean(call_gemini(PROMPT.format(n=n + 2, level=level, doc=text)), n)
    if len(result["questions"]) < 3:
        raise UserError("Couldn't build enough valid questions from this document. Try again.", 502)
    result["truncated"] = truncated
    result["chars"] = len(text)
    return jsonify(result)


@app.errorhandler(UserError)
def on_user_error(e):
    return jsonify(error=str(e)), e.status


@app.errorhandler(413)
def too_large(_):
    return jsonify(error=f"File is too large. Maximum is {MAX_UPLOAD_MB} MB."), 413


@app.errorhandler(Exception)
def on_error(e):
    if hasattr(e, "code") and isinstance(e.code, int) and e.code < 500:
        return jsonify(error=getattr(e, "description", "Bad request")), e.code
    app.logger.exception(e)
    return jsonify(error="Something went wrong on the server."), 500


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", 5000)), debug=False)
